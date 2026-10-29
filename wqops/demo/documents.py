"""Fictional quality documents of the demo company (ADR-0014 item 9): procedures, work instructions, forms, records
and certificates, in the formats companies really keep (Word, PDF, Excel, CSV, text and a scanned image).

Problems are built in on purpose so the library has something to find: an older revision of a procedure, an
identical copy of a form, a work instruction without a number, and processes with no documents at all (complaints,
training procedure). Everything is invented; the files are made by code, never copied from a real company.

The folder is in the company's language (ADR-0023, step 23e): ``files("zh")`` gives the Chinese demo company's folder
(Chinese names and content, its suppliers and people, metric units), ``files()`` in English is Harbor Lane's, the same
bytes as before. Document numbers, item codes and gauge codes are the same in both.
"""

from __future__ import annotations

import io
import json
import os
import sys
import zipfile

from wqops import i18n as I18N
from wqops.demo import scenario as S
from wqops.i18n import say

ROOT = "Harbor Lane QMS"                      # the US company's; files(lang) and load() use the company's own
LABEL = "Demo: Harbor Lane QMS folder"
UPLOADER = "dana@demo.example"
COMPANY = "Harbor Lane Supply (Demo)"


def names(lang: str | None = None) -> dict[str, str]:
    """The folder's root, the import's label and the company name in a language ('en' gives the constants above)."""
    with I18N.use(lang):
        return {"root": say("Harbor Lane QMS"), "label": say("Demo: Harbor Lane QMS folder"),
                "company": say("Harbor Lane Supply (Demo)")}


# ---------------------------------------------------------------- file makers

def _latin1(s: str) -> bool:
    try:
        s.encode("latin-1")
    except UnicodeEncodeError:
        return False
    return True


def _wrap(line: str, width: float = 44.0) -> list[str]:
    """Break a line with Chinese so it fits the page (a Chinese character is 1 em, a Latin one about half)."""
    out, cur, w = [], "", 0.0
    for ch in line:
        cw = 0.5 if ord(ch) < 0x2000 else 1.0
        if w + cw > width and cur:
            out.append(cur)
            cur, w = "", 0.0
        cur += ch
        w += cw
    return [*out, cur]


def pdf(lines: list[str]) -> bytes:
    """A plain text PDF (Helvetica, 46 lines a page) without any library.

    A document with text Helvetica cannot show (Chinese) is written with the standard Chinese font STSong-Light
    (Adobe-GB1, not embedded: every PDF reader has it), its text as UTF-16 code units (encoding UniGB-UCS2-H), so the
    text can be read back and searched. A document in Latin letters only stays exactly as before."""
    def esc(s: str) -> str:
        return s.encode("latin-1", "replace").decode("latin-1").replace("\\", "\\\\").replace("(", "\\(") \
            .replace(")", "\\)")

    cjk = not all(_latin1(ln) for ln in lines)
    if cjk:
        lines = [part for ln in lines for part in _wrap(ln)]
    pages = [lines[i:i + 46] for i in range(0, max(1, len(lines)), 46)] or [[]]
    objs: list[bytes] = []
    n_pages = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    font = 3 + 2 * n_pages
    for i, page in enumerate(pages):
        if cjk:
            text = " ".join(f"<{ln.encode('utf-16-be').hex().upper()}> '" for ln in page)
        else:
            text = " ".join(f"({esc(ln)}) '" for ln in page)
        stream = "BT /F1 11 Tf 56 790 Td 15 TL " + text + " ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents {4 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode())
        data = stream.encode("latin-1")
        objs.append(b"<< /Length " + str(len(data)).encode() + b" >>\nstream\n" + data + b"\nendstream")
    if cjk:
        objs.append(f"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light /Encoding /UniGB-UCS2-H "
                    f"/DescendantFonts [{font + 1} 0 R] >>".encode())
        objs.append(f"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light /CIDSystemInfo "
                    f"<< /Registry (Adobe) /Ordering (GB1) /Supplement 4 >> /FontDescriptor {font + 2} 0 R "
                    f"/DW 1000 /W [1 95 500] >>".encode())
        objs.append(b"<< /Type /FontDescriptor /FontName /STSong-Light /Flags 6 /FontBBox [-25 -254 1000 880] "
                    b"/ItalicAngle 0 /Ascent 880 /Descent -120 /CapHeight 880 /StemV 93 >>")
    else:
        objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def docx(header: list[str], paragraphs: list[str], table: list[list[str]] | None = None, *,
         author: str = COMPANY) -> bytes:
    import docx as D  # noqa: PLC0415 - python-docx, a hub dependency

    d = D.Document()
    d.core_properties.author = author
    hdr = d.sections[0].header
    hdr.paragraphs[0].text = header[0] if header else ""
    for h in header[1:]:
        hdr.add_paragraph(h)
    for p in paragraphs:
        if p.startswith("# "):
            d.add_heading(p[2:], level=1)
        elif p.startswith("## "):
            d.add_heading(p[3:], level=2)
        else:
            d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
    buf = io.BytesIO()
    d.save(buf)
    return _stable_zip(buf.getvalue())


def xlsx(sheets: dict[str, list[list]], *, author: str = COMPANY) -> bytes:
    import openpyxl  # noqa: PLC0415

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.properties.creator = author
    buf = io.BytesIO()
    wb.save(buf)
    return _stable_zip(buf.getvalue())


def _stable_zip(data: bytes) -> bytes:
    """Office files are zips with timestamps; fix them so the same document gives the same bytes every time."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                import re  # noqa: PLC0415

                body = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-01-05T09:00:00Z", body)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 5, 9, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, body)
    return out.getvalue()


def scan(lines: list[str]) -> bytes:
    """A 'scanned' certificate: dark text on an off-white page, as a PNG, read back with OCR."""
    from PIL import Image, ImageDraw, ImageFont  # noqa: PLC0415

    img = Image.new("L", (1700, 1100), 246)
    draw = ImageDraw.Draw(img)
    big, small = ImageFont.load_default(size=46), ImageFont.load_default(size=34)
    y = 90
    for i, ln in enumerate(lines):
        f = big if i == 0 else small
        draw.text((110, y), ln, fill=25, font=f)
        y += 78 if i == 0 else 58
    draw.rectangle((60, 50, 1640, 1050), outline=80, width=4)
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    return buf.getvalue()


# ---------------------------------------------------------------- content

def _bare(name: str) -> str:
    """A demo company's name without its demo mark: 'Gulf Fasteners', '海湾紧固件'."""
    return name.removesuffix(" (Demo)").removesuffix("（演示）")


def _people(lang: str) -> dict[str, str]:
    """Role -> name of the demo people who appear in the documents: Harbor Lane's in English, else the active
    company's (scenario.activate(manifest))."""
    people, names_ = (S.EN["PEOPLE"], S.EN["NAMES"]) if lang == "en" else (S.PEOPLE, S.NAMES)
    return {role: names_.get(email, email) for role, email in people.items()}


def _procedure(co: str, no: str, title: str, rev: str, eff: str, owner: str, purpose: str,  # noqa: PLR0917
               steps: list[str], records: list[str]) -> tuple[list[str], list[str]]:
    header = [say("{company} - Quality procedure", company=co),
              say("Doc No.: {no}    Rev: {rev}    Effective date: {eff}", no=no, rev=rev, eff=eff)]
    body = [f"# {title}", say("Document owner: {owner}", owner=owner), say("## 1. Purpose"), purpose,
            say("## 2. Scope"), say("This procedure applies to the Tampa Distribution Center and the Head Office."),
            say("## 3. Procedure"), *[f"3.{i} {s}" for i, s in enumerate(steps, 1)], say("## 4. Records"),
            *[f"- {r}" for r in records], say("## 5. Revision history"),
            say("Rev {rev}: current issue, effective {eff}.", rev=rev, eff=eff)]
    return header, body


def _rev_name(no: str, title: str, rev: str) -> str:
    return say("{no} {title} Rev {rev}", no=no, title=title, rev=rev)


def files(lang: str | None = None) -> list[tuple[str, bytes]]:
    """(path inside the import, bytes) for every demo document, deterministic; in ``lang`` (default: the current
    language). The Chinese folder names the active company's people (scenario.activate)."""
    lang = I18N.norm(lang) or I18N.current()
    with I18N.use(lang):
        return _files(lang)


def _files(lang: str) -> list[tuple[str, bytes]]:  # noqa: PLR0915 - one folder, file after file
    zh = lang == "zh"
    nm = names(lang)
    root, co = nm["root"], nm["company"]
    prof = S.ZH if zh else S.EN
    gulf, bayline, coastal, summit = prof["GULF"], prof["BAYLINE"], prof["COASTAL"], prof["SUMMIT"]
    item = {i[0]: i for i in prof["ITEMS"]}
    who = _people(lang)
    owner_n, ops_n, buyer_n, qa_n, wh_n = (who.get(r, r) for r in ("owner", "ops_manager", "purchasing", "quality",
                                                                    "warehouse"))
    lab = say("Sunshine Metrology (Demo)")
    qm = say("Quality Manager")

    def doc(header, paragraphs, table=None):
        return docx(header, paragraphs, table, author=co)

    def book(sheets):
        return xlsx(sheets, author=co)

    out: list[tuple[str, bytes]] = []
    P = f"{root}/{say('01 Procedures')}"
    f001, f002 = say("F-QA-001 Document Change Request"), say("F-QA-002 Master Document List")
    f010, f012 = say("F-QA-010 Receiving Inspection Record"), say("F-QA-012 Nonconformance Report")
    fpu1, fpu2 = say("F-PU-001 Approved Supplier List"), say("F-PU-002 Supplier Evaluation")

    h, b = _procedure(co, "QP-01", say("Control of Documents"), "C", "2025-02-03", qm,
                      say("To make sure the right revision of every document is used and obsolete ones are not."),
                      [say("Documents are reviewed and approved by the Quality Manager before use."),
                       say("Changes are requested on the Document Change Request F-QA-001."),
                       say("The Master Document List F-QA-002 shows the current revision of every document."),
                       say("Obsolete documents are removed from points of use or marked OBSOLETE.")],
                      [f001, f002])
    out.append((f"{P}/{_rev_name('QP-01', say('Control of Documents'), 'C')}.docx", doc(h, b)))

    title = say("Purchasing and Supplier Evaluation")
    h, b = _procedure(co, "QP-04", title, "B", "2025-04-14", say("Purchasing Lead"),
                      say("To buy only from suppliers able to meet our requirements and to watch how they perform."),
                      [say("New suppliers are evaluated on the Supplier Evaluation form F-PU-002 before the first "
                           "order."),
                       say("Approved suppliers are listed on the Approved Supplier List F-PU-001."),
                       say("Purchase orders state item, quantity, price and required date."),
                       say("On-time delivery and rejected lots are reviewed for each supplier every quarter.")],
                      [fpu1, fpu2])
    out.append((f"{P}/{_rev_name('QP-04', title, 'B')}.docx", doc(h, b)))

    title = say("Receiving Inspection")
    h, b = _procedure(co, "QP-05", title, "D", "2025-09-08", qm,
                      say("To check incoming goods against the order and the inspection plan before they are "
                          "stocked."),
                      [say("The receiving clerk checks quantity and labels against the purchase order."),
                       say("Items with an inspection plan are sampled and measured; results go on the Receiving "
                           "Inspection Record F-QA-010."),
                       say("Rejected lots are moved to QC Hold and a Nonconformance Report F-QA-012 is opened."),
                       say("Accepted lots are moved to Main Stock.")],
                      [f010, f012])
    out.append((f"{P}/{_rev_name('QP-05', title, 'D')}.pdf", pdf(h + b)))

    title8 = say("Control of Nonconforming Product")

    def qp08(rev: str, eff: str, extra: list[str]) -> bytes:
        h, b = _procedure(co, "QP-08", title8, rev, eff, qm,
                          say("To keep nonconforming product from being used or shipped by mistake."),
                          [say("Nonconforming product is tagged and moved to QC Hold."),
                           say("The finder opens a Nonconformance Report F-QA-012."),
                           say("Quality decides the disposition: use as is, rework, return to supplier or scrap."),
                           *extra],
                          [f012])
        return pdf(h + b)

    out.append((f"{P}/{_rev_name('QP-08', title8, 'C')}.pdf",
                qp08("C", "2025-06-02", [say("Major nonconformances need a corrective action under QP-09."),
                                         say("Use-as-is needs approval by a second person."),
                                         say("The lot number of the affected product is recorded on the NCR before "
                                             "the disposition."),
                                         say("The General Manager signs off every major NCR before it is closed.")])))
    out.append((f"{P}/{say('folder::old')}/{_rev_name('QP-08', title8, 'B')}.pdf",
                qp08("B", "2023-11-20", [say("Repeated problems are discussed in the weekly meeting.")])))

    title = say("Corrective and Preventive Action")
    f014 = say("F-QA-014 CAPA Request")
    h, b = _procedure(co, "QP-09", title, "B", "2025-06-02", qm,
                      say("To find and remove the causes of problems so they do not happen again."),
                      [say("A corrective action is requested on the CAPA Request F-QA-014."),
                       say("The owner finds the root cause (5 Whys or 8D) and plans actions with due dates."),
                       say("Quality checks that the actions worked before the CAPA is closed.")],
                      [f014])
    out.append((f"{P}/{_rev_name('QP-09', title, 'B')}.docx", doc(h, b)))

    f020 = say("F-QA-020 Calibration Log")
    h, b = _procedure(co, "QP-11", say("Calibration of Measuring Equipment"), "A", "2024-08-19", qm,
                      say("To make sure gauges used for receiving inspection measure correctly."),
                      [say("Every gauge has an ID and is listed on the Calibration Log F-QA-020."),
                       say("Gauges are calibrated every 12 months by an accredited laboratory."),
                       say("A gauge found out of tolerance is withdrawn and the lots it measured are reviewed.")],
                      [f020])
    out.append((f"{P}/{_rev_name('QP-11', say('file title::Calibration'), 'A')}.docx", doc(h, b)))

    W = f"{root}/{say('02 Work Instructions')}"
    title = say("Picking and Packing")
    out.append((f"{W}/WI-WH-03 {title}.docx", doc(
        [say("{company} - Work instruction", company=co),
         say("Doc No.: {no}    Rev: {rev}    Effective date: {eff}", no="WI-WH-03", rev="2", eff="2025-05-12")],
        [f"# {title}", say("Owner: {owner}", owner=say("Warehouse Lead")),
         say("1. Print the pick list from ERPNext for the sales order."),
         say("2. Pick from Main Stock only; never from QC Hold."),
         say("3. Pack in the box size on the pick list and seal with packing tape."),
         say("4. Print the 4 x 6 shipping label and scan it to confirm the shipment.")])))
    out.append((f"{W}/{say('torque check for anchors')}.txt", "\n".join([   # deliberately without a number
        say("Torque check for wedge anchors"), "",
        say("Use the torque wrench (gauge ID TW-02)."),
        say("Pick 3 anchors from each new lot of ANCHOR-38."),
        say("Set the test block and torque to 25 ft-lb."),
        say("Write the result on the receiving inspection record."), ""]).encode("utf-8")))

    Fm = f"{root}/{say('03 Forms')}"
    out.append((f"{Fm}/{f010}.xlsx", book({say("sheet::Form"): [
        [say("Receiving Inspection Record"), None, None, None, say("Form No. {no}", no="F-QA-010"),
         say("Rev {rev}", rev="3")],
        [say("Date"), say("PO / Receipt"), say("Item"), say("Sample size"), say("Result (accept / reject)"),
         say("Inspector")],
        *[[None] * 6 for _ in range(3)]]})))
    ncr_form = doc([co, say("Form No. {no}    Rev {rev}", no="F-QA-012", rev="4")],
                   [say("# Nonconformance Report"), say("NCR number: ________   Date: ________"),
                    say("Item / lot: ____________________   Quantity: ______"),
                    say("Description of the nonconformance:"), "", say("Disposition (circle one): use as is / rework / "
                                                                       "return to supplier / scrap"),
                    say("Approved by: ____________   Date: ________")])
    out.append((f"{Fm}/{f012}.docx", ncr_form))
    out.append((f"{Fm}/{say('copy of {name}', name=f012)}.docx", ncr_form))     # an identical copy
    out.append((f"{Fm}/{fpu1}.xlsx", book({say("sheet::Approved suppliers"): [
        [say("Approved Supplier List"), None, None, say("Form No. {no}", no="F-PU-001"), say("Rev {rev}", rev="5")],
        [say("Supplier"), say("Supplies"), say("Approved on"), say("Certificate"), say("Status")],
        [gulf, item["HX-BOLT-M8"][2], "2024-03-11", "ISO 9001", say("Approved - on watch (late deliveries)")],
        [bayline, item["CBL-12-50"][2], "2023-09-05", "ISO 9001", say("supplier status::Approved")],
        [coastal, item["BOX-12"][2], "2024-01-22", "-", say("supplier status::Approved")],
        [summit, item["TAPE-MEAS-25"][2], "2025-02-17", "ISO 9001", say("supplier status::Approved")],
        [lab, say("Calibration services"), "2024-11-04", "ISO/IEC 17025", say("supplier status::Approved")]]})))

    rec = say("04 Records")
    R = f"{root}/{rec}/2026"
    passed = say("calibration result::Pass")
    out.append((f"{R}/{say('F-QA-020 Calibration Log 2026')}.xlsx", book({say("sheet::Log"): [
        [say("Calibration Log"), None, None, None, say("Form No. {no}", no="F-QA-020"), say("Rev {rev}", rev="1")],
        [say("Gauge ID"), say("Description"), say("Calibrated on"), say("Due"), say("Result"), say("Lab")],
        ["CAL-01", say("Thread pitch gauge M8"), "2026-01-14", "2027-01-14", passed, lab],
        ["CAL-02", say("Digital caliper 0-150 mm"), "2026-01-14", "2027-01-14", passed, lab],
        ["TW-02", say("Torque wrench 10-80 ft-lb"), "2025-07-03", "2026-07-03", passed, lab]]})))
    out.append((f"{R}/{say('Receiving inspection {item} lot {lot}', item='HX-BOLT-M8', lot='4471')}.pdf", pdf([
        say("Receiving Inspection Record    Form No. F-QA-010 Rev 3"), "",
        say("Date: {date}    Receipt: {receipt}    Item: {item}", date="2026-08-12", receipt="MAT-PRE-DEMO-0041",
            item="HX-BOLT-M8"), say("Supplier: {supplier}", supplier=gulf),
        say("Sample size: 8    Characteristic: pitch diameter 7.160 - 7.212 mm"),
        say("Readings: {values}", values="7.190 7.214 7.188 7.216 7.195 7.213 7.192 7.191"),
        say("Result: REJECT (3 of 8 above the upper limit)"), say("Inspector: {name}", name=qa_n),
        say("Lot moved to QC Hold; NCR opened.")])))
    out.append((f"{R}/{say('Internal audit report {month}', month='2026-03')}.docx", doc(
        [say("{company} - Internal audit report", company=co)],
        [say("# Internal Audit Report - Receiving and Storage"),
         say("Audit date: {date}    Auditor: {name}", date="2026-03-18", name=ops_n),
         say("Scope: QP-05 Receiving Inspection and QP-08 Control of Nonconforming Product."),
         say("Finding 1 (minor): two receiving records did not show the sample size."),
         say("Finding 2 (observation): QC Hold area label is faded."),
         say("Follow-up: CAPA requested for finding 1 on F-QA-014.")])))
    attendees = ("、" if zh else ", ").join([owner_n, ops_n, qa_n, buyer_n])
    out.append((f"{R}/{say('Management review minutes {month}', month='2026-06')}.docx", doc(
        [say("{company} - Management review", company=co)],
        [say("# Management Review Minutes"), say("Date: {date}    Chair: {name}", date="2026-06-24", name=owner_n),
         say("Attendees: {names}", names=attendees),
         say("Inputs reviewed: on-time delivery, supplier performance, nonconformances, audit results."),
         say("Decision: keep {supplier} on watch; review again in Q4.", supplier=gulf),
         say("Decision: start a document library and move all procedures into it.")])))

    L = f"{root}/{rec}/{say('folder::Logs')}"
    minor, major, closed = say("ncr class::Minor"), say("ncr class::Major"), say("log status::Closed")
    rts = say("Return to supplier")
    out.append((f"{L}/{say('NCR log {years}', years='2025-2026')}.xlsx", book({say("sheet::NCR log"): [
        [say("Nonconformance log - {company}", company=co)],
        [say("NCR #"), say("Date"), say("Item"), say("Supplier"), say("Qty"), say("Description"), say("Class"),
         say("Disposition"), say("Status"), say("log column::Closed")],
        ["NCR-2025-01", "2025-02-11", "HX-BOLT-M8", gulf, 200, say("Thread damage on 14 bolts in box 3"),
         minor, rts, closed, "2025-02-20"],
        ["NCR-2025-02", "2025-04-03", "CBL-12-50", bayline, 5, say("Insulation nicked on 2 coils"),
         minor, say("log disposition::Scrap"), closed, "2025-04-05"],
        ["NCR-2025-03", "2025-06-17", "BOX-18", coastal, 50, say("Boxes crushed in delivery"), minor,
         say("log disposition::Use as is"), closed, "2025-06-18"],
        ["NCR-2025-04", "2025-09-09", "HX-BOLT-M8", gulf, 400,
         say("Pitch diameter above limit on 3 of 8 samples"), major, rts, closed, "2025-10-02"],
        # another hand's way of writing: an abbreviation and a date in another format
        ["NCR-2026-01", "2026-01-22", "LEVEL-24", summit, 12, say("Bubble vial misaligned"), minor, say("RTV"),
         closed, "2026/1/30" if zh else "1/30/2026"],
        ["NCR-2026-02", "2026-03-05", "TAPE-48", coastal, 30,
         say("Tape does not stick at low temperature"), minor, None, say("log status::Open"), None],
        ["NCR-2026-03", "2026-05-14", "WRAP-20", None, None, say("Stretch wrap roll damaged in warehouse"), minor,
         say("log disposition::Scrap"), say("Cancelled - entered twice"), None]]})))
    out.append((f"{L}/{say('CAPA log')}.xlsx", book({say("sheet::CAPA"): [
        [say("CAR No."), say("Opened"), say("Type"), say("Problem"), say("Root cause"), say("Action"), say("Owner"),
         say("NCR Ref"), say("Status"), say("log column::Closed")],
        ["CAR-2025-01", "2025-09-12", say("capa type::Corrective"),
         say("{supplier} bolts out of tolerance", supplier=_bare(gulf)),
         say("Supplier thread rolling die worn"), say("Supplier replaced the die; pitch check added to QP-05"), qa_n,
         "NCR-2025-04", closed, "2025-12-15"],
        ["CAR-2026-01", "2026-02-02", say("capa type::Corrective"),
         say("Late deliveries from {supplier}", supplier=_bare(gulf)),
         say("No safety stock for the key item"), say("Reorder point set; weekly call with the supplier"), buyer_n,
         None, say("log status::In progress"), None],
        ["PAR-2026-01", "2026-04-20", say("capa type::Preventive"), say("Tape adhesion in cold storage"), None,
         say("Test tape samples before the winter purchase"), qa_n, "NCR-2026-02", say("log status::Open"),
         None]]})))
    competent = say("training result::Competent")
    rows = [say("Name,Course,Date,Trainer,Result"),
            ",".join([wh_n, say("WI-WH-03 Picking and Packing"), "2025-05-20", ops_n, competent]),
            ",".join([wh_n, say("QP-05 Receiving Inspection (receiving clerk part)"), "2025-09-15", qa_n,
                      competent]),
            ",".join([buyer_n, say("QP-04 Purchasing and Supplier Evaluation"), "2025-04-28", owner_n, competent])]
    # UTF-8 (with a byte order mark for Chinese, so Excel opens it right; the hub reads both)
    out.append((f"{root}/{rec}/{say('folder::Training')}/{say('training records warehouse')}.csv",
                ("\n".join(rows) + "\n").encode("utf-8-sig" if zh else "utf-8")))

    C = f"{root}/{say('05 Certificates')}"
    # The scanned registrar's certificate stays in English for both companies (registrars' certificates often are,
    # and the scan's plain font has no Chinese); the Chinese supplier is named in pinyin there.
    cert_name = "Binhai Electrical (Demo)" if zh else bayline
    out.append((f"{C}/{say('{supplier} ISO 9001 certificate', supplier=_bare(bayline))}.png", scan([
        "CERTIFICATE OF REGISTRATION (FICTIONAL)", "This certifies that the quality management system of",
        cert_name, "has been assessed and found to conform to ISO 9001:2015",
        "Scope: distribution of electrical wiring products", "Certificate No. DEMO-QMS-20417",
        "Valid until: 2027-09-30", "Issued by: Example Registrar (Demo)"])))
    tape = item["TAPE-MEAS-25"][1]
    out.append((f"{C}/{say('CoC {supplier} lot {lot}', supplier=_bare(summit), lot='2208')}.pdf", pdf([
        say("CERTIFICATE OF CONFORMANCE"), summit, "",
        say("Customer: {company}    PO: {po}", company=co, po="PUR-ORD-DEMO-0117"),
        say("Item: {item} ({code})", item=tape, code="TAPE-MEAS-25"),
        say("Lot: {lot}    Quantity: {qty}", lot="2208", qty="120"),
        say("We certify that the goods above conform to the purchase order and to our"),
        say("product specification TM-25 Rev 6."),
        say("Signed: Quality Department, {supplier}    Date: {date}", supplier=summit, date="2026-07-09")])))
    return out


def load(conn, manifest: dict, today, model=None) -> dict:
    """Import the demo folder as the owner would, take its text and classify it (with the model when the company has
    a key, else by the rules). Nothing is confirmed: that is for people. Once only. The folder is in the company's
    language and is imported by the company's own owner (manifest people)."""
    from wqops.hub.docs import ai_classify as AIC  # noqa: PLC0415
    from wqops.hub.docs import library as LIB  # noqa: PLC0415

    if not manifest.get("demo"):
        raise ValueError(say("demo documents are only loaded into a company manifest marked demo: true"))
    S.activate(manifest)
    lang = I18N.norm((manifest.get("company") or {}).get("language")) or I18N.company_language()
    label = names(lang)["label"]
    uploader = S.PEOPLE.get("owner", UPLOADER)
    done = conn.execute("select id from doc_imports where label = %s", (label,)).fetchone()
    if done:
        return {"import": done["id"], "loaded": False}
    imp = LIB.new_import(conn, uploader, label)
    for path, data in files(lang):
        LIB.add_file(conn, imp["id"], path, data)
    with I18N.use(lang):
        LIB.extract_import(conn, imp["id"])
        out = AIC.run(conn, {"email": uploader, "roles": ["owner"]}, imp["id"], model, today=today,
                      company=manifest["company"]["name"])
    return {"import": imp["id"], "loaded": True, "files": out["files"], "by_model": out["by_model"],
            "by_rules": out["by_rules"]}


def main() -> int:
    """python -m wqops.demo.documents: add the demo documents to a demo installation loaded before they existed."""
    from wqops import company as C  # noqa: PLC0415
    from wqops.hub import db as DB  # noqa: PLC0415
    from wqops.hub.ai import llm as L  # noqa: PLC0415
    from wqops.hub.ont import api as OA  # noqa: PLC0415

    manifest = C.load(os.environ["WQO_COMPANY_FILE"])
    if not manifest.get("demo"):
        print("demo documents are only loaded into a company manifest marked demo: true", file=sys.stderr)
        return 1
    I18N.set_company_language(manifest["company"].get("language"))
    with DB.connect(os.environ["WQO_DB"]) as conn:
        DB.migrate(conn)
        print(json.dumps(load(conn, manifest, OA.today(manifest), L.from_env()), indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
