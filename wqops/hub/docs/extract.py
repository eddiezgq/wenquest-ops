"""Text from imported files (ADR-0014 item 2): local code only, no network. PDF, Word, Excel, CSV and text are read
directly; scanned pages and images go through local OCR (Tesseract, English). A file that yields no text is marked
``no_text`` and can still be classified by a person.
"""

from __future__ import annotations

import csv
import io
import mimetypes
import re
from dataclasses import dataclass
from pathlib import PurePosixPath

from wqops.i18n import say

TEXT = {".txt", ".md", ".csv"}
IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
OLD = {".doc", ".xls"}                                  # kept, but their text cannot be read here
ACCEPTED = {".pdf", ".docx", ".xlsx", ".xlsm"} | TEXT | IMAGES | OLD
MAX_CHARS = 2_000_000
MAX_ROWS = 5000
MAX_OCR_PIXELS = 25_000_000                             # about 25 MB in grayscale
MIN_PDF_CHARS = 40                                      # fewer per page than this: treat the PDF as scanned


@dataclass
class Text:
    body: str
    status: str                                         # ok | ocr | no_text | error
    note: str = ""


def suffix(name: str) -> str:
    return PurePosixPath(name.lower()).suffix


def accepted(name: str) -> bool:
    return suffix(name) in ACCEPTED


def mime(name: str) -> str:
    return mimetypes.guess_type(name)[0] or "application/octet-stream"


def clean(s: str) -> str:
    s = s.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"[ \t ]+", " ", s)
    s = re.sub(r"\n[ ]+", "\n", s)
    return re.sub(r"\n{3,}", "\n\n", s).strip()[:MAX_CHARS]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def ocr_image(img) -> str:
    import pytesseract  # noqa: PLC0415 - only needed for scans

    w, h = img.size
    if min(w, h) < 20 or max(w, h) > 40 * min(w, h):     # rules, lines and slivers hold no text
        return ""
    img = img.convert("L")
    k = min(1000 / min(w, h), (MAX_OCR_PIXELS / (w * h)) ** 0.5, 4.0)
    if k > 1:                                           # small scans read better enlarged, within a memory cap
        img = img.resize((int(w * k), int(h * k)))
    return pytesseract.image_to_string(img, lang="eng")


def _ocr_ready() -> bool:
    try:
        import pytesseract  # noqa: PLC0415

        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _image(data: bytes) -> Text:
    from PIL import Image  # noqa: PLC0415

    if not _ocr_ready():
        return Text("", "no_text", say("OCR (Tesseract) is not installed on this server"))
    with Image.open(io.BytesIO(data)) as img:
        pages = []
        for i in range(getattr(img, "n_frames", 1)):
            img.seek(i)
            pages.append(ocr_image(img.copy()))
    body = clean("\n\n".join(pages))
    return Text(body, "ocr" if body else "no_text", "" if body else say("OCR found no text"))


def _pdf(data: bytes) -> Text:
    import pypdf  # noqa: PLC0415

    reader = pypdf.PdfReader(io.BytesIO(data))
    pages = [(p.extract_text() or "") for p in reader.pages]
    body = clean("\n\n".join(pages))
    if len(body) >= MIN_PDF_CHARS * max(1, len(pages)):
        return Text(body, "ok")
    if not _ocr_ready():
        return Text(body, "ok" if body else "no_text", "" if body else say("scanned PDF; OCR is not installed"))
    from PIL import Image  # noqa: PLC0415

    scanned = []
    for p in reader.pages:
        for im in p.images:
            with Image.open(io.BytesIO(im.data)) as img:
                scanned.append(ocr_image(img))
    ocr = clean("\n\n".join(scanned))
    if len(ocr) > len(body):
        return Text(ocr, "ocr")
    return Text(body, "ok" if body else "no_text", "" if body else say("no text found in the PDF"))


def _docx(data: bytes) -> Text:
    import docx  # noqa: PLC0415

    d = docx.Document(io.BytesIO(data))
    parts = []
    for s in d.sections:                                # "Doc No." and "Rev" often live in the header
        for p in s.header.paragraphs:
            parts.append(p.text)
        for t in s.header.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
    parts += [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
    for s in d.sections:
        parts += [p.text for p in s.footer.paragraphs]
    body = clean("\n".join(parts))
    return Text(body, "ok" if body else "no_text")


def _xlsx(data: bytes) -> Text:
    import openpyxl  # noqa: PLC0415

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"Sheet: {ws.title}")
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= MAX_ROWS:
                parts.append(f"(more rows not read; first {MAX_ROWS} kept)")
                break
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                parts.append(" | ".join(cells).rstrip(" |"))
    wb.close()
    body = clean("\n".join(parts))
    return Text(body, "ok" if body else "no_text")


def _csv(data: bytes) -> Text:
    rows = list(csv.reader(io.StringIO(_decode(data))))[:MAX_ROWS]
    body = clean("\n".join(" | ".join(r) for r in rows))
    return Text(body, "ok" if body else "no_text")


def _plain(data: bytes) -> Text:
    body = clean(_decode(data))
    return Text(body, "ok" if body else "no_text")


READERS = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".csv": _csv, ".txt": _plain, ".md": _plain}


def text_of(name: str, data: bytes) -> Text:
    """The text of one file, by its extension. Never raises: a broken file is ``error`` with a note."""
    ext = suffix(name)
    if ext in OLD:
        return Text("", "no_text", say("old {ext} format: save it as {ext}x to get its text", ext=ext))
    reader = READERS.get(ext) or (_image if ext in IMAGES else None)
    if reader is None:
        return Text("", "no_text", say("format not read"))
    try:
        return reader(data)
    except Exception as exc:  # a damaged or unusual file must not stop the import
        return Text("", "error", say("could not read the file: {error}", error=type(exc).__name__))
