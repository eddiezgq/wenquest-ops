"""The model drafts a document a control point is missing (ADR-0016, step 16c), such as the quality policy or a
procedure for customer property. The draft goes into the document library as a file to review: a person confirms
it, sends it for approval and someone else approves it (ADR-0014). The model never makes a document effective, and
anything it does not know about the company it leaves as "[to fill: ...]" for a person. Without a model key this is
not offered.
"""

from __future__ import annotations

import datetime as dt
import io
import re

from wqops import i18n as I18N
from wqops.hub.ai import agent as AG
from wqops.hub.ai import llm as L
from wqops.hub.ai import tools as T
from wqops.hub.docs import library as LIB
from wqops.hub.system import qms as QMS
from wqops.hub.system import rules as R
from wqops.i18n import say

TYPES = ("procedure", "work_instruction", "form", "other")
MAX_CHARS = 20000
PROCESS = {"production": "production", "traceability": "production", "customer_property": "production",
           "design": "production", "post_delivery": "complaints", "changes": "management_review",
           "scope": "management_review", "processes": "management_review", "policy": "management_review",
           "roles": "management_review", "ncr_concession": "nonconformance", "capa_for_major": "capa",
           "capa_on_time": "capa",
           # ISO 13485 (17a)
           "regulatory_role": "management_review", "outsourced": "purchasing", "qms_software": "document_control",
           "device_file": "document_control", "records": "document_control", "management_rep": "management_review",
           "communication": "management_review", "infrastructure": "production", "work_environment": "production",
           "contamination": "production", "planning": "production", "risk_management": "other",
           "customer_requirements": "other", "customer_communication": "complaints", "purchasing_info": "purchasing",
           "cleanliness": "production", "installation": "production", "servicing": "production",
           "sterile_records": "production", "process_validation": "production", "sterilization_validation":
           "production", "identification": "production", "implantable": "production", "preservation": "production",
           "process_monitoring": "management_review", "product_release": "production",
           "after_delivery": "nonconformance", "rework": "nonconformance", "data_analysis": "management_review",
           "preventive": "capa"}


class NoModel(RuntimeError):
    pass


def draftable(conn, standard: str) -> dict[str, dict]:
    """Control points shown by a document that have none in effect yet."""
    v = QMS.view(conn, standard)
    out = {}
    for p in v["processes"]:
        for c in p["controls"]:
            r = c.get("result") or {}
            if (r.get("rule") or c["rule"]) in R.LINKED and r.get("status") in (R.UNKNOWN, R.NOT_MET):
                out[c["id"]] = c
    return out


def tool() -> dict:
    return {
        "name": "write_document",
        "description": "Write the document as a draft for the company to review.",
        "input_schema": {"type": "object", "properties": {
            "title": {"type": "string", "description": "the document's title, e.g. 'Quality Policy'"},
            "doc_type": {"type": "string", "enum": list(TYPES)},
            "sections": {"type": "array", "items": {"type": "object", "properties": {
                "heading": {"type": "string"},
                "paragraphs": {"type": "array", "items": {"type": "string"}}}, "required": ["heading", "paragraphs"]}},
            "to_fill": {"type": "array", "items": {"type": "string"},
                        "description": "what a person must fill in or decide before approval"}},
            "required": ["title", "doc_type", "sections"]},
    }


def _docx(company: str, title: str, sections: list[dict], to_fill: list[str]) -> bytes:
    import docx as D  # noqa: PLC0415 - python-docx, a hub dependency

    d = D.Document()
    hdr = d.sections[0].header
    hdr.paragraphs[0].text = say("{company} - DRAFT written by AI for review; not in effect until approved",
                                 company=company)
    d.add_heading(title, level=1)
    for s in sections:
        d.add_heading(s["heading"], level=2)
        for para in s["paragraphs"]:
            d.add_paragraph(para)
    if to_fill:
        d.add_heading(say("To fill in before approval"), level=2)
        for t in to_fill:
            d.add_paragraph(t, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _t(v, n: int) -> str:
    return CONTROL_CHARS.sub("", str(v if v is not None else "")).strip()[:n]


def _items(v) -> list:
    return v if isinstance(v, list) else ([v] if v not in (None, "") else [])


def _clean(inp: dict) -> dict:
    title = re.sub(r"\s+", " ", _t(inp.get("title"), 200))[:120]
    sections = []
    for s in _items(inp.get("sections")):
        if not isinstance(s, dict):
            continue
        paras = [_t(x, 3000) for x in _items(s.get("paragraphs")) if _t(x, 3000)]
        if _t(s.get("heading"), 120) and paras:
            sections.append({"heading": _t(s["heading"], 120), "paragraphs": paras[:40]})
    if not title or not sections:
        raise L.LlmError(say("the model's draft has no title or no text"))
    if sum(len(p) for s in sections for p in s["paragraphs"]) > MAX_CHARS:
        raise L.LlmError(say("the model's draft is too long"))
    return {"title": title, "doc_type": inp.get("doc_type") if inp.get("doc_type") in TYPES else "other",
            "sections": sections[:20], "to_fill": [_t(x, 300) for x in _items(inp.get("to_fill")) if _t(x, 300)][:30]}


def draft(conn, user: dict, standard: str, control: str, model: L.Anthropic | None, *, today: dt.date,
          company: str = "the company") -> dict:
    """Write the draft and put it in the library for review. Returns the new document."""
    if model is None or not model.configured():
        raise NoModel(say("there is no model key; write the document yourself and link it"))
    c = draftable(conn, standard).get(control)
    if c is None:
        raise QMS.QmsError(say("this control point is not shown by a document, or it already has one in effect"))
    have = conn.execute("select number, title, doc_type from documents where status in ('effective', 'confirmed', "
                        "'draft') and title is not null order by number nulls last limit 80").fetchall()
    known = "\n".join(f"- {d['number'] or ''} {d['title']} ({d['doc_type']})".strip() for d in have) or "none yet"
    std = QMS.name_of(standard)
    ctx = T.Ctx(conn=conn, user=user, today=today)
    AG.start(ctx, "document", say("draft for {standard} {control}: {requirement}", standard=standard, control=control,
                                  requirement=I18N.label(c["summary"]))[:1000])
    out = {"engine": "model", "model": model.model, "tool_calls": [], "verified": True, "unverified": [], "usage": {}}
    try:
        with I18N.use(I18N.company_language()):        # a company document: written in the company's language
            resp = model.create(
                f"You help {company} set up its quality management system. Write a short, plain draft document "
                "that shows how the company meets one requirement. Keep it practical: purpose, scope, "
                "responsibilities, what is done, which records are kept. Never invent facts about the company "
                "(names, numbers, dates, products, sites): write [to fill: ...] where they belong and list them "
                "under to_fill. Refer to the company's existing documents by number where they fit. A person "
                "reviews and approves the draft; you never make it effective.",
                [{"role": "user", "content": f"Requirement ({std} {', '.join(c['clauses'])}): {c['summary']}\n"
                                             f"Usually shown by: {c.get('hint') or 'a procedure'}\n\n"
                                             f"The company's documents:\n{known}"}],
                [tool()], 4000, tool_choice={"type": "tool", "name": "write_document"})
        if resp.get("stop_reason") == "max_tokens":
            raise L.LlmError(say("the model's draft was cut off"))
        use = next((b for b in resp.get("content", []) if b.get("type") == "tool_use"), None)
        if use is None or not isinstance(use.get("input"), dict):
            raise L.LlmError(say("the model did not write a draft"))
        d = _clean(use["input"])
        try:
            with I18N.use(I18N.company_language()):      # a company document: the company's language
                data = _docx(company, d["title"], d["sections"], d["to_fill"])
        except ValueError as exc:
            raise L.LlmError(say("the model's draft could not be written as a file: {exc}", exc=exc)) from exc
        name = re.sub(r"[^\w .-]+", " ", f"DRAFT {d['title']}").strip()[:100] + ".docx"
        shown_std, clauses = I18N.label(std), ", ".join(c["clauses"])
        with conn.transaction():
            imp = LIB.new_import(conn, user["email"], say("AI draft for {standard} {clauses}", standard=shown_std,
                                                          clauses=clauses))
            doc = LIB.add_file(conn, imp["id"], name, data)
            LIB.extract_one(conn, LIB.get(conn, doc["id"]))
            conn.execute("update documents set ai_drafted = true where id = %s", (doc["id"],))
            row = LIB.propose(conn, doc["id"], {"doc_type": d["doc_type"], "process": PROCESS.get(control, "other"),
                                                "title": d["title"], "confidence": "low",
                                                "summary": say("AI draft for: {requirement}",
                                                               requirement=I18N.label(c["summary"]))},
                                 "ai", ctx.run_id)
            conn.execute("update doc_imports set status = 'done', engine = 'ai' where id = %s", (imp["id"],))
            LIB.log(conn, doc["id"], user["email"], "drafted by AI",
                    say("for {standard} {clauses}; {n} things to fill in", standard=shown_std, clauses=clauses,
                        n=len(d["to_fill"])))
        out.update(usage=resp.get("usage", {}), answer=d["title"],
                   tool_calls=[{"tool": "write_document", "args": {"control": control}, "ok": True,
                                "summary": say("{title}: {n} sections", title=d["title"],
                                               n=len(d["sections"]))}])
        AG.finish(ctx, out)
        return row
    except Exception as exc:                    # the run is closed whatever went wrong, then the error goes on
        AG.finish(ctx, out | {"engine": "model (failed)", "answer": "", "verified": False,
                              "error": f"{type(exc).__name__}: {exc}"[:1000]})
        raise
