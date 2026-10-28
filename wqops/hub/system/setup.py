"""Setting the quality system up for the company (ADR-0016, step 16d): a short questionnaire whose answers propose
what does not apply; changes proposed from the company's own procedures (by the model, system/setup_ai.py), each
quoting the procedure, checked here and decided by a person; and procedures written from the workflow in use, by
code from templates, so that what the procedure says is what the system does.

Plain code: no AI, no network (tests/test_isolation.py). Every change goes through the functions people use by hand
(qms.declare, qms.set_setting, qms.link), with the same checks and signatures.
"""

from __future__ import annotations

import datetime as dt
import io

from psycopg.types.json import Json

from wqops import i18n as I18N
from wqops.hub.docs import library as LIB
from wqops.hub.quality import records as Q
from wqops.hub.quality import workflow as W
from wqops.hub.quality import workflow_proposals as WP
from wqops.hub.system import qms as QMS
from wqops.hub.system import rules as R
from wqops.i18n import say

# ---------------------------------------------------------------- the questionnaire

def questions(conn, standard: str) -> dict:
    """The questionnaire of the standard's pack: the one adopted, or the pack file before it is adopted (or when the
    adopted pack is older than questionnaires in packs)."""
    row = QMS.enabled(conn, standard)
    return (row and row["content"].get("questions")) or QMS.load(standard).get("questions") or {}


class SetupError(ValueError):
    pass


def clean_answers(qs: dict, raw: dict) -> dict:
    if not isinstance(raw, dict):
        raise SetupError(say("answers are given by question"))
    out = {}
    for k, v in raw.items():
        q = qs.get(k)
        if q is None:
            raise SetupError(say("there is no question {k}", k=repr(k)))
        if v in (None, ""):
            continue
        if q["type"] == "yesno":
            if type(v) is not bool:
                raise SetupError(say("{label} yes or no", label=I18N.label(q['label'])))
            out[k] = v
        elif q["type"] == "number":
            if type(v) is not int or not 0 < v < 1_000_000:
                raise SetupError(say("{label}: a whole number", label=I18N.label(q['label'])))
            out[k] = v
        else:
            out[k] = str(v).strip()[:500]
    return out


def answers(conn, standard: str) -> dict:
    r = conn.execute("select answers, email, at from qms_setup where standard = %s order by id desc limit 1",
                     (standard,)).fetchone()
    return {"answers": r["answers"], "email": r["email"], "at": r["at"].isoformat()} if r else {"answers": {}}


def save_answers(conn, email: str, standard: str, raw: dict) -> dict:
    vals = clean_answers(questions(conn, standard), raw)
    conn.execute("insert into qms_setup (standard, answers, email) values (%s, %s, %s)", (standard, Json(vals), email))
    return vals


def proposed_applicability(conn, standard: str) -> list[dict]:
    """What the answers say does not apply (or does), against what is declared now. The reasons are in the reader's
    language: declared, they are kept as written."""
    vals = answers(conn, standard)["answers"]
    now = QMS.applicability(conn, standard)
    out = []
    for k, q in questions(conn, standard).items():
        if not q.get("if_no") or k not in vals:
            continue
        for control, reason in q["if_no"].items():
            applies = bool(vals[k])
            question = I18N.label(q["label"])
            why = I18N.label(reason) if not applies else \
                say("Answered yes: {question}.", question=question.rstrip("?？").lower())
            current = now.get(control)
            out.append({"control": control, "applies": applies, "reason": why, "question": question,
                        "same": bool(current) and current["applies"] == applies})
    return out


def apply_answers(conn, user: dict, standard: str, password: str, today: dt.date,  # noqa: PLR0917
                  tz: str = "UTC") -> int:
    """Declare, with the person's signature, what the answers say; only what differs from now."""
    todo = [p for p in proposed_applicability(conn, standard) if not p["same"]]
    if not todo:
        raise SetupError(say("the declarations already match the answers"))
    for p in todo:
        QMS.declare(conn, user, standard, p["control"], p["applies"], p["reason"], password, today, tz)
    return len(todo)


# ---------------------------------------------------------------- changes proposed from the procedures

KINDS = {"setting": "a number of days or one of our own rules switched on or off",
         "applies": "a requirement that does not apply (or does), with the reason",
         "evidence": "a procedure that shows how a requirement is met"}
MAX_ITEMS = 30


def clean(pack: dict, raw: dict) -> dict:
    if not isinstance(raw, dict) or not isinstance(raw.get("kind"), str) or raw["kind"] not in KINDS:
        raise SetupError(say("a change is one of: {KINDS}", KINDS=', '.join(KINDS)))
    k = raw["kind"]
    cs = QMS.controls(pack)
    if k == "setting":
        key = str(raw.get("key") or "")
        spec = (pack.get("settings") or {}).get(key)
        if spec is None:
            raise SetupError(say("there is no setting {key}", key=repr(key)))
        v = raw.get("value")
        if type(spec["value"]) is bool:
            if type(v) is not bool:
                raise SetupError(say("{label}: yes or no", label=I18N.label(spec['label'])))
        elif type(v) is not int or not spec.get("min", 0) <= v <= spec.get("max", 3650):
            raise SetupError(say("{label}: a whole number from {spec} to {spec2}",
                                 label=I18N.label(spec['label']), spec=spec.get('min', 0),
                                 spec2=spec.get('max', 3650)))
        return {"kind": k, "key": key, "value": v}
    control = str(raw.get("control") or "")
    if control not in cs:
        raise SetupError(say("there is no requirement {control}", control=repr(control)))
    if k == "applies":
        if not cs[control].get("may_exclude"):
            raise SetupError(say("this requirement always applies"))
        if type(raw.get("applies")) is not bool or not str(raw.get("reason") or "").strip():
            raise SetupError(say("say whether it applies and why"))
        return {"kind": k, "control": control, "applies": raw["applies"], "reason": str(raw["reason"]).strip()[:500]}
    if cs[control]["rule"] not in R.LINKED and not cs[control].get("switch"):
        raise SetupError(say("this requirement is judged from the records, not from a document"))
    return {"kind": k, "control": control}


def describe(pack: dict, ch: dict) -> str:
    """One proposed change in plain words; anything malformed reads as a change WenQuest Ops cannot make."""
    try:
        clean(pack, ch)
    except (SetupError, TypeError, ValueError, KeyError):
        return say("A change WenQuest Ops cannot make: {change}",
                   change=", ".join(f"{k} {v}" for k, v in ch.items())[:200])
    if ch["kind"] == "setting":
        spec = pack["settings"][ch["key"]]
        v = (say("setting::on") if ch["value"] else say("setting::off")) if type(ch["value"]) is bool \
            else ch["value"]
        return say("{setting}: {value}", setting=I18N.label(spec["label"]), value=v)
    c = QMS.controls(pack)[ch["control"]]
    req = I18N.label(c["summary"])
    if ch["kind"] == "applies":
        return say("Applies: {requirement}", requirement=req) if ch["applies"] else \
            say("Does not apply: {requirement}", requirement=req)
    return say("Evidence for: {requirement}", requirement=req)


def check(conn, standard: str, pack: dict, ch: dict, doc_id: int | None, quote: str,  # noqa: PLR0917
          body: str | None) -> dict:
    problems = []
    verified = body is not None and WP.quoted(quote, body)
    if body is None:
        problems.append(say("it does not name one of the procedures that were read"))
    elif not verified:
        problems.append(say("its quote is not in the procedure it names"))
    already = False
    if ch["kind"] == "setting":
        already = QMS.settings(conn, standard, pack).get(ch["key"]) == ch["value"]
    elif ch["kind"] == "applies":
        now = QMS.applicability(conn, standard).get(ch["control"])
        already = bool(now) and now["applies"] == ch["applies"]
    else:
        already = doc_id in {d["id"] for d in QMS.evidence(conn, standard).get(ch["control"], [])}
    return {"verified": verified, "problems": problems, "already": already}


def store(conn, email: str, standard: str, docs: list[dict], result: dict,  # noqa: PLR0917
          ai_run: int | None) -> dict:
    row = QMS.enabled(conn, standard)
    if row is None:
        raise QMS.QmsError(say("adopt the standard first"))
    pack = row["content"]
    bodies = {d["id"]: d["text"] for d in docs}
    changes = result.get("changes") if isinstance(result.get("changes"), list) else []
    with conn.transaction():
        p = conn.execute("insert into qms_setup_proposals (standard, documents, created_by, ai_run, summary) values "
                         "(%s, %s, %s, %s, %s) returning *",
                         (standard, Json([{k: d[k] for k in ("id", "name", "revision", "status")} for d in docs]),
                          email, ai_run, str(result.get("summary") or "")[:2000])).fetchone()
        for seq, raw in enumerate(changes[:MAX_ITEMS], 1):
            item = raw if isinstance(raw, dict) else {}
            doc = item.get("document_id") if type(item.get("document_id")) is int else None
            quote = str(item.get("quote") or "").strip()[:1000]
            try:
                ch = clean(pack, item.get("change"))
                found = check(conn, standard, pack, ch, doc, quote, bodies.get(doc))
            except (SetupError, TypeError, ValueError) as exc:
                ch = {k: str(v)[:200] for k, v in item["change"].items()} if isinstance(item.get("change"), dict) \
                    else {}
                found = {"verified": doc in bodies and WP.quoted(quote, bodies[doc]), "problems": [str(exc)],
                         "already": False}
            conn.execute("insert into qms_setup_items (proposal_id, seq, change, document_id, quote, why, verified, "
                         "problems, already) values (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
                         (p["id"], seq, Json(ch), doc if doc in bodies else None, quote,
                          str(item.get("why") or "")[:1000], found["verified"], Json(found["problems"]),
                          found["already"]))
    return p


def _iso(r: dict) -> dict:
    return {k: (v.isoformat() if hasattr(v, "isoformat") else v) for k, v in r.items()}


def acceptable(it: dict) -> bool:
    return it["decision"] == "open" and it["verified"] and not it["problems"] and not it["already"]


def view(conn, pid: int) -> dict:
    p = conn.execute("select * from qms_setup_proposals where id = %s", (pid,)).fetchone()
    if p is None:
        raise KeyError(pid)
    row = QMS.enabled(conn, p["standard"])
    pack = row["content"] if row else QMS.load(p["standard"])
    docs = {d["id"]: d for d in p["documents"]}
    items = conn.execute("select * from qms_setup_items where proposal_id = %s order by seq", (pid,)).fetchall()
    return {**_iso(p), "items": [{**_iso(i), "text": describe(pack, i["change"]),
                                  "document": docs.get(i["document_id"]),
                                  "acceptable": acceptable(i)} for i in items]}


def listing(conn, standard: str) -> list[dict]:
    rows = conn.execute("select p.id, p.created_at, p.created_by, p.summary, p.documents, count(i.id) as changes, "
                        "count(i.id) filter (where i.decision = 'accepted') as accepted from qms_setup_proposals p "
                        "left join qms_setup_items i on i.proposal_id = p.id where p.standard = %s group by p.id "
                        "order by p.id desc limit 20", (standard,)).fetchall()
    return [_iso(r) for r in rows]


def decide(conn, user: dict, iid: int, accept: bool, *, password: str = "", note: str = "",
           today: dt.date, tz: str = "UTC") -> dict:
    """Accept (make the change, as a person would by hand) or reject one proposed change. The item is locked while
    the change is made, so it is decided and made once."""
    if accept:
        it = conn.execute("select i.*, p.standard from qms_setup_items i join qms_setup_proposals p on "
                          "p.id = i.proposal_id where i.id = %s", (iid,)).fetchone()
        if it is not None and acceptable(it) and it["change"].get("kind") == "applies":
            Q.check_signer(conn, user, password)       # before the transaction: a failed attempt stays recorded
    with conn.transaction():
        it = conn.execute("select i.*, p.standard from qms_setup_items i join qms_setup_proposals p on "
                          "p.id = i.proposal_id where i.id = %s for update of i", (iid,)).fetchone()
        if it is None:
            raise KeyError(iid)
        if it["decision"] != "open":
            raise SetupError(say("this change was already accepted by {decided_by}", decided_by=it["decided_by"])
                             if it["decision"] == "accepted" else
                             say("this change was already rejected by {decided_by}", decided_by=it["decided_by"]))
        if accept:
            if not acceptable(it):
                raise SetupError(say("only a change whose quote was found and that passes the checks can be accepted"))
            ch, std = it["change"], it["standard"]
            if ch["kind"] == "setting":
                QMS.set_setting(conn, user["email"], std, ch["key"], ch["value"], today, tz)
            elif ch["kind"] == "applies":
                QMS.declare(conn, user, std, ch["control"], ch["applies"], ch["reason"], password, today, tz)
            else:
                QMS.link(conn, user["email"], std, ch["control"], it["document_id"], True, today,
                         say("proposed from the procedure: \"{quote}\"", quote=it["quote"][:300]), tz)
        return conn.execute("update qms_setup_items set decision = %s, decided_by = %s, decided_at = now(), "
                            "decision_note = %s where id = %s returning *",
                            ("accepted" if accept else "rejected", user["email"], note.strip()[:1000] or None,
                             iid)).fetchone()


# ---------------------------------------------------------------- procedures written from the workflow in use

TITLES = {"ncr": ("Control of Nonconforming Product", "nonconformance"),
          "capa": ("Corrective and Preventive Action", "capa")}


def _title(kind: str) -> str:
    return {"ncr": lambda: say("Control of Nonconforming Product"),
            "capa": lambda: say("Corrective and Preventive Action")}[kind]()


def _zh() -> bool:
    return I18N.current() == "zh"


def _list(items) -> str:
    return ("、" if _zh() else ", ").join(items)


def _or(items) -> str:
    return (" 或 " if _zh() else " or ").join(items)


def _approvals(cfg: dict, step: str) -> list[str]:
    L, out = I18N.label, []
    for a in cfg["approvals"]:
        if a["before"] != step:
            continue
        roles = _or(L(r) for r in a["roles"])
        when = [say("when {field} is {values}", field=L(k), values=_or(L(y) for y in x))
                for k, x in (a.get("when") or {}).items()]
        out.append(say("{approval} (signed by {roles}, {conditions})", approval=L(a["label"]), roles=roles,
                       conditions=_list(when)) if when
                   else say("{approval} (signed by {roles})", approval=L(a["label"]), roles=roles))
    return out


def procedure_text(conn, kind: str) -> tuple[str, list[tuple[str, list[str]]]]:
    """The procedure for a kind of record, written from the enabled workflow: (title, [(heading, paragraphs)]), in
    the current language (the workflow's own names are translated where the catalog has them)."""
    v = W.ensure(conn, kind)
    cfg = W.view(v["config"])
    L = I18N.label
    title = _title(kind)
    label = {s: L(x.get("label", s)) for s, x in cfg["states"].items()}
    open_states = [label[s] for s, x in cfg["states"].items() if not x.get("final")]
    steps = []
    for name, st in cfg["steps"].items():
        step, frm = L(st.get("label") or name), _list(label[s] for s in st["from"])
        bits = [say("{step} (from {source} to {target})", step=step, source=frm, target=label[st["to"]])
                if st.get("to") else say("{step} (from {source})", step=step, source=frm)]
        bits.append(say("Done by: {roles}.", roles=_list(L(r.replace("_", " ")) for r in st["who"])))
        if st["needs"]:
            bits.append(say("Needs: {list}.", list=_list(L(x) for x in st["needs"])))
        if st.get("reason"):
            bits.append(say("A reason is recorded."))
        if st.get("sign"):
            bits.append(say("Signed with the person's password: \"{sign}\".", sign=L(st["sign"])))
        if st["not_by_text"]:
            bits.append(say("Not done by {list}.", list=_or(st["not_by_text"])))
        own = [L(f["label"]) for f in cfg["fields"].values() if name in (f.get("required_at") or [])]
        if own:
            bits.append(say("Recorded first: {list}.", list=_list(own)))
        appr = _approvals(cfg, name)
        if appr:
            bits.append(say("Approved first: {list}.", list=("；" if _zh() else "; ").join(appr)))
        steps.append(("" if _zh() else " ").join(bits))
    rules = []
    for sev in (cfg.get("rules") or {}).get("capa_closed_before_close") or []:
        rules.append(say("A {severity} nonconformance is closed only after its corrective action is closed.",
                         severity=L(sev)))
    for sev in (cfg.get("rules") or {}).get("eight_d_before_close") or []:
        rules.append(say("A {severity} nonconformance is closed only after its 8D report is complete.",
                         severity=L(sev)))
    if (cfg.get("rules") or {}).get("one_open_capa_per_ncr"):
        rules.append(say("A nonconformance has at most one open corrective action at a time."))
    num = cfg["numbering"]
    name_of = cfg.get("label") or kind
    one = L(name_of) if _zh() else name_of.lower()
    many = L(name_of) if _zh() else f"{name_of.lower()}s"
    finals = _or(label[s] for s, x in cfg["states"].items() if x.get("final"))
    sections = [
        (say("1. Purpose"), [say("To set out how {records} are recorded, decided and closed, as WenQuest Ops carries "
                                 "them out. This procedure is written from the workflow in use; when the workflow "
                                 "changes, it is written again.", records=many)]),
        (say("2. Scope"), [say("Every {record} recorded in WenQuest Ops.", record=one)]),
        (say("3. Records"), [say("Each is numbered {first} onwards and moves through: {states}, then {final}.",
                                 first=f"{num['prefix']}-{'0' * (int(num['digits']) - 1)}1",
                                 states=_list(open_states), final=finals),
                             say("Every step is logged with who did it and when; signatures and the log cannot be "
                                 "changed or deleted; a closed or cancelled record is locked.")]),
        (say("4. Steps"), steps),
    ]
    if cfg["fields"]:
        sections.append((say("5. Information the company records"), [
            say("{field} ({type})", field=L(f["label"]), type=f["type"]) for f in cfg["fields"].values()]))
    if rules:
        sections.append((say("{n}. Rules", n=len(sections) + 1), rules))
    sections.append((say("{n}. Source", n=len(sections) + 1), [
        say("{workflow} workflow version {version}, enabled by {by} on {date}, fingerprint {fp}.",
            workflow=L(name_of), version=v["version"], by=v["enabled_by"], date=f"{v['enabled_at']:%Y-%m-%d}",
            fp=v["fingerprint"][:12])]))
    return title, sections


def _docx(company: str, title: str, sections: list[tuple[str, list[str]]]) -> bytes:
    import docx as D  # noqa: PLC0415 - python-docx, a hub dependency

    d = D.Document()
    d.sections[0].header.paragraphs[0].text = say("{company} - written by WenQuest Ops from the workflow in use; not "
                                                  "in effect until approved", company=company)
    d.add_heading(title, level=1)
    for h, paras in sections:
        d.add_heading(h, level=2)
        for p in paras:
            d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def generate(conn, email: str, kind: str, company: str = "the company") -> dict:
    """Write the procedure for a kind of record into the library, to review and have approved. A company document:
    written in the company's language."""
    if kind not in TITLES:
        raise KeyError(kind)
    v = W.ensure(conn, kind)
    source = f"{kind} workflow v{v['version']}"          # stored and compared (generated): not translated
    with I18N.use(I18N.company_language()):
        title, sections = procedure_text(conn, kind)
        data = _docx(company, title, sections)
        file_name = say("{title} (from {source})", title=title, source=source) + ".docx"
    summary = say("Written by WenQuest Ops from the {source}.", source=source)
    note = say("Written from the {source}", source=source)
    with conn.transaction():
        imp = LIB.new_import(conn, email, note)
        doc = LIB.add_file(conn, imp["id"], file_name, data)
        LIB.extract_one(conn, LIB.get(conn, doc["id"]))
        conn.execute("update documents set generated_from = %s where id = %s", (source, doc["id"]))
        row = LIB.propose(conn, doc["id"], {"doc_type": "procedure", "process": TITLES[kind][1], "title": title,
                                            "revision": f"W{v['version']}", "confidence": "high",
                                            "summary": summary}, "generated")
        conn.execute("update doc_imports set status = 'done', engine = 'rules' where id = %s", (imp["id"],))
        LIB.log(conn, doc["id"], email, "written from the workflow", source)
    return row


def generated(conn) -> list[dict]:
    """Procedures written from workflows, and whether the workflow has changed since."""
    current = {k: f"{k} workflow v{W.ensure(conn, k)['version']}" for k in TITLES}
    rows = conn.execute("select id, title, status, generated_from, created_at from documents where generated_from "
                        "is not null and status not in ('rejected', 'obsolete') order by id desc").fetchall()
    return [{**_iso(r), "current": r["generated_from"] == current.get(r["generated_from"].split(" ")[0])}
            for r in rows]
