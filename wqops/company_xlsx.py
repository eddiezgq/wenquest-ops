"""Excel template and import for the company manifest (docs/plan.md, Phase 1 item 2).

The template is the demo company filled in; a business edits it and imports it. Errors point at the sheet, row
and column ("People row 4, Email: ...") so a non-programmer can fix them.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from wqops import company as C

YES, NO = "yes", "no"

# Company sheet: (manifest path, label, help)
COMPANY_ROWS = [
    ("demo", "Demo company", "yes = fictional data for trying the platform; no = your real business"),
    ("company.name", "Company name", "The company created in ERPNext"),
    ("company.abbr", "Abbreviation", "Letters and digits, up to 6; used as a suffix on warehouse names"),
    ("company.country", "Country", "ERPNext country name, e.g. United States"),
    ("company.currency", "Currency", "ISO code, e.g. USD"),
    ("company.timezone", "Time zone", "e.g. America/New_York, America/Chicago, America/Los_Angeles"),
    ("company.fiscal_year_start", "Fiscal year starts", "MM-DD, e.g. 01-01"),
    ("company.date_format", "Date format", " / ".join(C.DATE_FORMATS)),
    ("company.units.weight", "Weight unit", " / ".join(C.WEIGHT_UNITS)),
    ("company.units.length", "Length unit", " / ".join(C.LENGTH_UNITS)),
]
# List sheets: (column label, key, help)
SITE_COLS = [("ID", "id", "lowercase letters, digits and -"), ("Name", "name", ""),
             ("Kind", "kind", " / ".join(C.SITE_KINDS)), ("Address line 1", "address.line1", "optional"),
             ("City", "address.city", ""), ("State", "address.state", ""),
             ("ZIP / postal code", "address.postal_code", ""), ("Country", "address.country", "")]
WH_COLS = [("ID", "id", "lowercase letters, digits and -"), ("Name", "name", ""),
           ("Site ID", "site", "an ID from the Sites sheet"), ("Kind", "kind", " / ".join(C.WAREHOUSE_KINDS))]
DEPT_COLS = [("ID", "id", "lowercase letters, digits and -"), ("Name", "name", "")]
PEOPLE_COLS = [("Name", "name", ""), ("Email", "email", "used to sign in"),
               ("Roles", "roles", "comma separated: " + ", ".join(C.ROLES)),
               ("Department ID", "department", "optional; an ID from the Departments sheet")]
LISTS = [("Sites", "sites", SITE_COLS), ("Warehouses", "warehouses", WH_COLS),
         ("Departments", "departments", DEPT_COLS), ("People", "people", PEOPLE_COLS)]

_HEAD = Font(bold=True)
_FILL = PatternFill("solid", fgColor="DDEBF7")


def _get(d: dict, path: str):
    for k in path.split("."):
        if not isinstance(d, dict):
            return None
        d = d.get(k)
    return d


def _set(d: dict, path: str, value) -> None:
    ks = path.split(".")
    for k in ks[:-1]:
        d = d.setdefault(k, {})
    d[ks[-1]] = value


def write_template(path, m: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    for line in [
        "WenQuest Ops - company manifest",
        "Fill in each sheet, then run: python -m wqops.company import this.xlsx company.yaml",
        "Rows that are completely empty are ignored. Errors name the sheet, row and column.",
        "This template is filled with a FICTIONAL demo company - replace it with your own business.",
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 100

    ws = wb.create_sheet("Company")
    ws.append(["Field", "Value", "Help"])
    for key, label, hint in COMPANY_ROWS:
        v = _get(m, key)
        if key == "demo":
            v = YES if v else NO
        ws.append([label, v, hint])
    _style(ws, [24, 36, 70])

    for title, key, cols in LISTS:
        ws = wb.create_sheet(title)
        ws.append([c[0] for c in cols])
        ws.append([c[2] for c in cols])
        for item in m.get(key) or []:
            row = []
            for _, k, _ in cols:
                v = _get(item, k)
                row.append(", ".join(v) if isinstance(v, list) else v)
            ws.append(row)
        _style(ws, [22] * len(cols), help_row=True)

    ws = wb.create_sheet("Modules")
    ws.append(["Module", "Turn on", "Help"])
    helps = {"core": "always on", "hub": "always on", "ai": "AI assistant", "quality": "quality management",
             "warehousing": "warehousing and distribution (Phase 2)", "production": "production (later)"}
    for mod in C.MODULES:
        ws.append([mod, YES if mod in m.get("modules", []) else NO, helps[mod]])
    _style(ws, [16, 12, 50])
    wb.save(path)


def _style(ws, widths, help_row=False):
    for cell in ws[1]:
        cell.font = _HEAD
        cell.fill = _FILL
    if help_row:
        for cell in ws[2]:
            cell.font = Font(italic=True, color="808080")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w


def _cell(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def read(path) -> dict:  # noqa: PLR0912, PLR0915 - straightforward sheet-by-sheet reading
    """Read a filled template into a manifest; raise CompanyError with sheet/row/column locations."""
    try:
        wb = load_workbook(Path(path), data_only=True)
    except Exception as exc:
        raise C.CompanyError([f"{path}: cannot open as an Excel workbook ({exc})"]) from exc
    errs: list[str] = []
    m: dict = {"version": C.SCHEMA_VERSION}
    where: dict[str, str] = {}

    if "Company" not in wb.sheetnames:
        raise C.CompanyError(["Company sheet is missing"])
    labels = {label: (p, r) for r, (p, label, _) in enumerate(COMPANY_ROWS)}
    for r, row in enumerate(wb["Company"].iter_rows(min_row=2, values_only=True), start=2):
        if not row or _cell(row[0]) is None:
            continue
        label = _cell(row[0])
        if label not in labels:
            errs.append(f"Company row {r}: unknown field {label!r}")
            continue
        p = labels[label][0]
        v = _cell(row[1]) if len(row) > 1 else None
        where[p] = f"Company row {r}, {label}"
        if p == "demo":
            if v not in (YES, NO, None):
                errs.append(f"Company row {r}, {label}: yes or no")
            v = v == YES
        if v is not None:
            _set(m, p, v)

    for title, key, cols in LISTS:
        items = []
        if title not in wb.sheetnames:
            errs.append(f"{title} sheet is missing")
            continue
        for r, row in enumerate(wb[title].iter_rows(min_row=3, values_only=True), start=3):
            vals = [_cell(v) for v in (row or ())][: len(cols)]
            if not any(vals):
                continue
            item: dict = {}
            for (_label, k, _), raw in zip(cols, vals + [None] * (len(cols) - len(vals)), strict=True):
                if raw is None:
                    continue
                _set(item, k, [x.strip() for x in raw.split(",") if x.strip()] if k == "roles" else raw)
            idx = len(items)
            for label, k, _ in cols:
                where[f"{key}[{idx}].{k}"] = f"{title} row {r}, {label}"
            where[f"{key}[{idx}]"] = f"{title} row {r}"
            items.append(item)
        m[key] = items

    if "Modules" in wb.sheetnames:
        mods = []
        for r, row in enumerate(wb["Modules"].iter_rows(min_row=2, values_only=True), start=2):
            if not row or _cell(row[0]) is None:
                continue
            on = _cell(row[1]) if len(row) > 1 else None
            if on not in (YES, NO):
                errs.append(f"Modules row {r}, Turn on: yes or no")
            elif on == YES:
                mods.append(_cell(row[0]))
            where["modules"] = "Modules sheet"
        m["modules"] = mods
    else:
        errs.append("Modules sheet is missing")

    for e in C.validate(m):
        loc, _, msg = e.partition(": ")
        errs.append(f"{_locate(loc, where)}: {msg}")
    if errs:
        raise C.CompanyError(errs)
    return C.normalize(m)


def _locate(loc: str, where: dict[str, str]) -> str:
    """Translate a manifest location (people[3].email) to a sheet location, falling back to the nearest parent."""
    probe = loc
    while probe:
        if probe in where:
            return where[probe]
        if "." in probe:
            probe = probe.rsplit(".", 1)[0]
        else:
            break
    return loc
